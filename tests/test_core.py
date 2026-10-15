import csv
import json
import sqlite3
import zipfile

import pytest
from openpyxl import load_workbook

from receipt_ocr_csv.core.common import ReceiptError, cents, settings
from receipt_ocr_csv.core.extraction import extract_fields, normalized_date, parse_qr


def boxes(*lines):
    return [
        {
            "page": 1,
            "bbox": [10, index * 30, 500, index * 30 + 20],
            "confidence": 0.95,
            "text": line,
        }
        for index, line in enumerate(lines)
    ]


def test_exact_money_and_date():
    assert cents("1,234.56") == 123456
    assert cents("0.10") + cents("0.20") == 30
    assert cents("-1.50") == -150
    for value in ["NaN", "Infinity", "1.009", True, "abc", "10000000000"]:
        with pytest.raises(ReceiptError):
            cents(value)
    assert normalized_date("开票日期：2026年2月3日") == "2026-02-03"
    assert normalized_date("2026-02-30") is None


def test_chinese_fields_qr_and_categories():
    fields, warnings = extract_fields(
        boxes(
            "销售方名称：示例酒店A",
            "购买方：示例公司A",
            "开票日期：2026年1月2日",
            "发票代码：123456789012",
            "发票号码：12345678",
            "税额：6.00",
            "价税合计：人民币 106.00",
            "校验码：12345678901234567890",
            "支付方式：银行卡",
        ),
        [],
    )
    assert fields["amount"]["value"] == "106.00"
    assert fields["tax"]["value"] == "6.00"
    assert fields["vendor"]["value"] == "示例酒店A"
    assert fields["category"]["value"] == "Lodging"
    assert fields["invoice_no"]["value"] == "12345678"
    assert fields["invoice_code"]["value"] == "123456789012"
    assert fields["check_code"]["value"] == "12345678901234567890"
    assert not warnings
    assert parse_qr("01,10,123456789012,12345678,100.00,20260102,123456") == {
        "invoice_code": "123456789012",
        "invoice_no": "12345678",
        "net_amount": "100.00",
        "date": "2026-01-02",
        "check_code": "123456",
    }
    qr = {"payload": "RECEIPT|date=2026-01-03|amount=120.00", "page": 1, "bbox": [0, 0, 20, 20]}
    fields, warnings = extract_fields(boxes("Date: 2026-01-02", "TOTAL: CNY 106.00"), [qr])
    assert fields["amount"]["confidence"] == 0.4
    assert len([w for w in warnings if "conflicts" in w]) == 2
    vat = {**qr, "payload": "01,10,123456789012,12345678,100.00,20260102,123456"}
    _, warnings = extract_fields(boxes("Date: 2026-01-02", "Tax: 6.00", "TOTAL: CNY 106.00"), [vat])
    assert not any("conflicts" in w for w in warnings)
    assert parse_qr("https://example.invalid") == {}


def test_missing_does_not_guess():
    fields, warnings = extract_fields(
        boxes("Subtotal: 100.00", "Net total: 101.00", "Cash tendered: 200.00"), []
    )
    assert fields["amount"]["value"] == ""
    assert any("amount" in warning for warning in warnings)


def test_text_pdf_bypasses_ocr_and_repeat_preserves_edits(store, text_pdf, monkeypatch):
    def forbidden(*args):
        raise AssertionError("OCR must not run on text-layer PDF")

    monkeypatch.setattr("receipt_ocr_csv.core.imaging.recognize", forbidden)
    result = store.scan(str(text_pdf))
    assert not result["errors"]
    identifier = result["processed"][0]["id"]
    item = store.show(identifier)
    assert item["engine"] == "pdf-text"
    assert store.list()[0]["file"].endswith("source.pdf")
    assert item["record"]["amount"] == "177.50"
    source = item["fields"]["amount"]["sources"][0]
    assert source["page"] == 1 and source["bbox"][0] >= 0
    assert store.preview(identifier, 1).is_file()
    edited = store.update(
        identifier, {"amount": "128.50", "category": "Transport", "note": "Reviewed sample"}
    )
    assert edited["status"] == "reviewed" and edited["fields"]["amount"]["edited"]
    again = store.scan(str(text_pdf))
    assert again["processed"][0]["reused"]
    assert store.show(identifier)["record"]["amount"] == "128.50"
    assert (
        len(
            store.list(
                {
                    "status": "reviewed",
                    "category": "Transport",
                    "date_from": "2026-01-01",
                    "date_to": "2026-01-31",
                }
            )
        )
        == 1
    )
    assert not store.list({"status": "parsed"})
    assert not store.list({"category": "Other"})
    assert not store.list({"date_to": "2026-01-01"})
    with pytest.raises(ReceiptError):
        store.preview(identifier, 2)
    with pytest.raises(ReceiptError):
        store.show("missing")
    for values in [
        {"amount": "1.001"},
        {"date": "2026-02-30"},
        {"currency": "us"},
        {"unknown": "x"},
        {"amount": True},
    ]:
        with pytest.raises((ReceiptError, ValueError)):
            store.update(identifier, values)


def test_exports_imports_backup_duplicates(store, text_pdf, tmp_path):
    identifier = store.scan(str(text_pdf))["processed"][0]["id"]
    store.update(identifier, {"amount": "128.50", "vendor": "=literal example"})
    for format in ["csv", "json", "xlsx"]:
        out = tmp_path / f"records.{format}"
        store.export(format, str(out))
        with pytest.raises(ReceiptError):
            store.export(format, str(out))
        store.export(format, str(out), True)
    assert (tmp_path / "records.csv").read_bytes().startswith(b"\xef\xbb\xbf")
    with (tmp_path / "records.csv").open(encoding="utf-8-sig", newline="") as file:
        assert next(csv.DictReader(file))["amount"] == "128.50"
    workbook = load_workbook(tmp_path / "records.xlsx", data_only=False)
    assert workbook.active.cell(2, 6).data_type == "s"
    imported = store.import_file(str(tmp_path / "records.json"))
    assert imported["imported"] == 1
    assert store.import_file(str(tmp_path / "records.csv"))["reused"] == 1
    assert len(store.list({"duplicates": True})) == 1
    store.update(imported["ids"][0], {"invoice_no": "OTHER0001"})
    assert len(store.list({"duplicates": True})) == 0
    backup = store.backup()
    with zipfile.ZipFile(backup["path"]) as archive:
        assert "receipts.db" in archive.namelist()
        assert any(name.endswith("page-1.png") for name in archive.namelist())
        archive.extract("receipts.db", tmp_path / "restore")
    with sqlite3.connect(tmp_path / "restore/receipts.db") as db:
        assert db.execute("SELECT COUNT(*) FROM records").fetchone()[0] == 2
    with pytest.raises(ReceiptError):
        store.backup(backup["path"])


def test_atomic_import_and_validation(store, tmp_path):
    path = tmp_path / "input.json"
    path.write_text(json.dumps([{"date": "2026-01-02", "amount": "12.30"}, {"amount": "1.234"}]))
    with pytest.raises(ReceiptError):
        store.import_file(str(path))
    assert store.list() == []
    path.write_text(json.dumps([{"amount": "12.30", "unexpected": 1}]))
    with pytest.raises(ReceiptError):
        store.import_file(str(path))
    for filters in [{"status": "bad"}, {"nope": 1}]:
        with pytest.raises(ReceiptError):
            store.list(filters)


def test_input_limits_and_symlink(store, text_pdf, tmp_path):
    link = tmp_path / "linked.pdf"
    link.symlink_to(text_pdf)
    with pytest.raises(ReceiptError):
        store.scan(str(link))
    folder = tmp_path / "folder"
    folder.mkdir()
    (folder / "corrupt.png").write_bytes(b"not an image")
    assert store.scan(str(folder))["errors"]
    with pytest.raises(ReceiptError):
        store.scan(str(tmp_path / "missing"))
    store.config["max_file_mb"] = 0
    assert store.scan(str(text_pdf))["errors"][0]["code"] == "FILE_TOO_LARGE"
    with pytest.raises(ReceiptError):
        store.scan(str(tmp_path / "data"))


def test_config(tmp_path, monkeypatch):
    monkeypatch.setenv("ROCR_DATA_DIR", str(tmp_path))
    assert settings()[0] == tmp_path
    (tmp_path / "config.yaml").write_text("max_pages: 2\n")
    assert settings()[1]["max_pages"] == 2
    (tmp_path / "config.yaml").write_text("max_pages: 0\n")
    with pytest.raises(ReceiptError):
        settings()
    (tmp_path / "config.yaml").write_text("unknown: true\n")
    with pytest.raises(ReceiptError):
        settings()
